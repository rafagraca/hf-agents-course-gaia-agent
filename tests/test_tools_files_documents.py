"""Tests for ``read_file`` on documents (PDF, Word, PowerPoint) and on media files.

Every file is generated inside the test (invented content only) and lives in the
temporary DATA folder, so nothing touches the network or the real DATA folder.
"""

from __future__ import annotations

import zipfile
from pathlib import Path

import pytest
from docx import Document
from pypdf import PdfReader, PdfWriter
from pypdf.errors import DependencyError

from gaia_agent.config import Settings
from gaia_agent.tools import _readers as readers
from gaia_agent.tools.files import ReadFileTool


@pytest.fixture
def attachments(settings: Settings) -> Path:
    """The folder where question attachments live (inside DATA)."""
    return settings.files_dir


@pytest.fixture
def read_file(settings: Settings) -> ReadFileTool:
    return ReadFileTool(settings)


def write_pdf(path: Path, lines: list[str]) -> None:
    """Write a minimal, valid single-page PDF that shows ``lines`` in Helvetica."""
    text_ops = " 0 -16 Td ".join(f"({line}) Tj" for line in lines)
    stream = f"BT /F1 12 Tf 72 720 Td {text_ops} ET".encode("latin-1")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    body = b"%PDF-1.4\n"
    offsets = []
    for number, content in enumerate(objects, start=1):
        offsets.append(len(body))
        body += f"{number} 0 obj\n".encode() + content + b"\nendobj\n"
    xref_at = len(body)
    body += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    body += b"".join(f"{offset:010d} 00000 n \n".encode() for offset in offsets)
    body += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_at}\n%%EOF\n".encode()
    path.write_bytes(body)


def write_pptx(path: Path, slides: list[list[str]]) -> None:
    """Write a bare-bones .pptx (a zip of slide XML parts) with one text run per string."""
    ns = (
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
        'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"'
    )
    with zipfile.ZipFile(path, "w") as archive:
        for number, texts in enumerate(slides, start=1):
            runs = "".join(f"<a:p><a:r><a:t>{text}</a:t></a:r></a:p>" for text in texts)
            archive.writestr(f"ppt/slides/slide{number}.xml", f"<p:sld {ns}>{runs}</p:sld>")


# --------------------------------------------------------------------------- read_file: documents


def test_read_file_extracts_pdf_text_by_page(read_file: ReadFileTool, attachments: Path) -> None:
    write_pdf(attachments / "paper.pdf", ["First invented line", "Second invented line"])

    result = read_file("paper.pdf")

    assert "First invented line" in result
    assert "Second invented line" in result
    assert "page 1" in result.lower()


def test_read_file_says_when_a_pdf_has_no_text(read_file: ReadFileTool, attachments: Path) -> None:
    write_pdf(attachments / "scan.pdf", [])

    result = read_file("scan.pdf")

    assert "no extractable text" in result.lower()


def test_read_file_reports_a_corrupt_pdf(read_file: ReadFileTool, attachments: Path) -> None:
    (attachments / "bad.pdf").write_bytes(b"%PDF-1.4 garbage without structure")

    assert read_file("bad.pdf").startswith("Error:")


def test_read_file_reports_a_password_protected_pdf(read_file: ReadFileTool, attachments: Path) -> None:
    writer = PdfWriter()
    writer.add_blank_page(72, 72)
    writer.encrypt("a-user-password")
    with (attachments / "locked.pdf").open("wb") as handle:
        writer.write(handle)

    result = read_file("locked.pdf")

    assert result.startswith("Error:")
    assert "password protected" in result


def test_read_file_explains_a_pdf_that_needs_an_optional_package(
    read_file: ReadFileTool, attachments: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    write_pdf(attachments / "aes.pdf", ["Invented line"])

    class NeedsCryptography:
        is_encrypted = True

        def __init__(self, source: str) -> None:
            pass

        def decrypt(self, password: str) -> int:
            raise DependencyError("cryptography>=3.1 is required for AES algorithm")

    monkeypatch.setattr(readers, "PdfReader", NeedsCryptography)

    result = read_file("aes.pdf")

    assert result.startswith("Error:")
    assert "cryptography" in result


def test_read_file_notes_when_a_pdf_has_more_pages_than_it_reads(read_file: ReadFileTool, attachments: Path) -> None:
    write_pdf(attachments / "single.pdf", ["Repeated invented line"])
    page = PdfReader(str(attachments / "single.pdf")).pages[0]
    writer = PdfWriter()
    for _ in range(61):
        writer.add_page(page)
    with (attachments / "long.pdf").open("wb") as handle:
        writer.write(handle)

    result = read_file("long.pdf")

    assert "Repeated invented line" in result
    assert "only the first 60 of 61 pages" in result


def test_read_file_extracts_docx_paragraphs_and_tables(read_file: ReadFileTool, attachments: Path) -> None:
    document = Document()
    document.add_paragraph("Invented opening paragraph")
    document.add_paragraph("   ")  # blank paragraphs are skipped
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "head-a"
    table.cell(0, 1).text = "head-b"
    table.cell(1, 0).text = "cell-a"
    table.cell(1, 1).text = "cell-b"
    document.save(str(attachments / "memo.docx"))

    result = read_file("memo.docx")

    assert "Invented opening paragraph" in result
    assert "head-a | head-b" in result
    assert "cell-a | cell-b" in result


def test_read_file_reports_a_corrupt_docx(read_file: ReadFileTool, attachments: Path) -> None:
    (attachments / "bad.docx").write_bytes(b"not a docx")

    assert read_file("bad.docx").startswith("Error:")


def test_read_file_says_when_a_docx_has_no_text(read_file: ReadFileTool, attachments: Path) -> None:
    Document().save(str(attachments / "blank.docx"))

    assert read_file("blank.docx") == "(the document has no text)"


def test_read_file_says_when_a_pptx_has_no_slides(read_file: ReadFileTool, attachments: Path) -> None:
    with zipfile.ZipFile(attachments / "empty.pptx", "w") as archive:
        archive.writestr("docProps/app.xml", "<Properties/>")

    assert read_file("empty.pptx") == "(the presentation has no slides)"


def test_read_file_marks_a_slide_without_text(read_file: ReadFileTool, attachments: Path) -> None:
    write_pptx(attachments / "pictures.pptx", [[]])

    assert read_file("pictures.pptx") == "Slide 1: (no text)"


def test_read_file_extracts_pptx_text_slide_by_slide(read_file: ReadFileTool, attachments: Path) -> None:
    write_pptx(attachments / "deck.pptx", [["Title one", "Detail one"], ["Title two"]])

    result = read_file("deck.pptx")

    assert "Slide 1" in result
    assert "Title one | Detail one" in result
    assert "Slide 2" in result
    assert "Title two" in result


def test_read_file_reports_a_corrupt_pptx(read_file: ReadFileTool, attachments: Path) -> None:
    (attachments / "bad.pptx").write_bytes(b"not a zip")

    assert read_file("bad.pptx").startswith("Error:")


@pytest.mark.parametrize("name", ["bomb.docx", "bomb.xlsx", "bomb.xlsm", "bomb.pptx"])
def test_read_file_refuses_an_office_file_that_unpacks_to_far_more_than_it_weighs(
    read_file: ReadFileTool, attachments: Path, monkeypatch: pytest.MonkeyPatch, name: str
) -> None:
    """A zip bomb is a few kilobytes on disk and gigabytes in memory: the declared sizes are added up first."""
    monkeypatch.setattr(readers, "MAX_UNPACKED_BYTES", 10_000)
    with zipfile.ZipFile(attachments / name, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("part-one.xml", b"0" * 6_000)
        archive.writestr("part-two.xml", b"0" * 6_000)

    result = read_file(name)

    assert result.startswith("Error:")
    assert "would unpack to" in result
    assert (attachments / name).stat().st_size < 1_000  # tiny on disk, as a bomb is


def test_an_office_file_that_unpacks_to_a_reasonable_size_is_still_read(
    read_file: ReadFileTool, attachments: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(readers, "MAX_UNPACKED_BYTES", 10_000)
    write_pptx(attachments / "fine.pptx", [["Hello slide"]])

    assert "Hello slide" in read_file("fine.pptx")


def test_the_unpacked_size_limit_is_150_megabytes() -> None:
    assert readers.MAX_UNPACKED_BYTES == 150 * 1024 * 1024


# --------------------------------------------------------------------------- read_file: media


@pytest.mark.parametrize("name", ["photo.png", "photo.JPG", "scan.webp"])
def test_read_file_points_images_to_describe_image(
    read_file: ReadFileTool, attachments: Path, monkeypatch: pytest.MonkeyPatch, name: str
) -> None:
    (attachments / name).write_bytes(b"\x89PNG fake image bytes")
    monkeypatch.setenv("GROQ_API_KEY", "fake-groq-key-for-tests")

    result = read_file(name)

    assert "describe_image" in result
    assert "fake-groq-key-for-tests" not in result


def test_read_file_says_images_cannot_be_seen_without_a_groq_key(read_file: ReadFileTool, attachments: Path) -> None:
    (attachments / "photo.png").write_bytes(b"\x89PNG fake image bytes")

    result = read_file("photo.png")

    assert "describe_image" not in result
    assert "cannot" in result.lower()


@pytest.mark.parametrize("name", ["voice.mp3", "talk.WAV", "memo.m4a", "song.flac"])
def test_read_file_points_audio_to_transcribe_audio(read_file: ReadFileTool, attachments: Path, name: str) -> None:
    (attachments / name).write_bytes(b"fake audio bytes")

    assert "transcribe_audio" in read_file(name)


def test_read_file_points_video_soundtracks_to_transcribe_audio(read_file: ReadFileTool, attachments: Path) -> None:
    (attachments / "clip.mp4").write_bytes(b"fake video bytes")

    result = read_file("clip.mp4")

    assert "transcribe_audio" in result
    assert "video" in result.lower()


@pytest.mark.parametrize("name", ["voice.mp3", "photo.png", "clip.mp4"])
def test_the_hints_name_the_file_but_never_the_folder_it_is_in(
    read_file: ReadFileTool, attachments: Path, monkeypatch: pytest.MonkeyPatch, name: str
) -> None:
    (attachments / name).write_bytes(b"fake bytes")
    monkeypatch.setenv("GROQ_API_KEY", "fake-groq-key-for-tests")

    result = read_file(name)

    assert f'file_path="{name}"' in result
    assert str(attachments) not in result
    assert attachments.as_posix() not in result
